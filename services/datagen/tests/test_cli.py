"""CLI: the legacy-workbook command (F05-FR-07)."""

from pathlib import Path

import datagen.cli
import pytest
from datagen.cli import main
from datagen.executor import Databases
from datagen.generate import generate
from datagen.legacy_workbook import SHEETS
from datagen.params import load_params
from openpyxl import load_workbook
from r2r_core.profile import load_profile
from sqlalchemy.exc import OperationalError


def test_f05_fr07_legacy_workbook_command_writes_the_workbook_without_a_database(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    def unreachable(*_: object) -> dict[str, str]:
        raise OperationalError("connect", {}, Exception("no database"))

    monkeypatch.setattr(datagen.cli, "lot_numbers_from_db", unreachable)
    out = tmp_path / "nested" / "legacy_tracker.xlsx"
    assert main(["legacy-workbook", "--out", str(out)]) == 0
    assert load_workbook(out).sheetnames == list(SHEETS)
    text = capsys.readouterr().out
    assert "no generated ERP database reachable" in text
    assert "typed statuses disagree" in text


@pytest.mark.integration
def test_f05_fr07_legacy_workbook_command_uses_the_inspection_lot_numbers_of_the_erp_database(
    tmp_path: Path, source_databases: Databases, monkeypatch: pytest.MonkeyPatch
) -> None:
    generate(load_profile("site_a"), load_params(), 4242, source_databases)
    monkeypatch.setenv("ERP_SIM_DSN", source_databases.erp)
    out = tmp_path / "legacy_tracker.xlsx"
    assert main(["legacy-workbook", "--out", str(out)]) == 0
    sheet = load_workbook(out)["Tracker"]
    lots = [row[3] for row in sheet.iter_rows(min_row=4, values_only=True) if row[0]]
    assert lots and all(str(lot).isdigit() and len(str(lot)) == 8 for lot in lots)
