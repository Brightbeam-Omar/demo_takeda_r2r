"""T7: the legacy tracker workbook (F05-FR-07, F05-AC-05)."""

from pathlib import Path

import pytest
from datagen.legacy_workbook import SHEETS, build_workbook, check_workbook, normalise_status
from datagen.model import Plan
from datagen.oracle import oracle_rows
from datagen.params import Params
from datagen.stats import compute_stats
from openpyxl import load_workbook
from r2r_core.profile import SiteProfile


@pytest.fixture(scope="module")
def workbook_path(
    plan: Plan, params: Params, profile: SiteProfile, tmp_path_factory: pytest.TempPathFactory
) -> Path:
    path = tmp_path_factory.mktemp("legacy") / "legacy_tracker.xlsx"
    build_workbook(plan, params, compute_stats(plan, profile)).save(path)
    return path


def test_f05_ac05_the_workbook_opens_and_has_the_twelve_tabs(workbook_path: Path) -> None:
    workbook = load_workbook(workbook_path)
    assert workbook.sheetnames == list(SHEETS)
    assert len(workbook.sheetnames) >= 12


def test_f05_fr07_it_is_deliberately_messy(workbook_path: Path) -> None:
    sheet = load_workbook(workbook_path)["Tracker"]
    assert sheet.merged_cells.ranges  # merged headers
    fills = {
        cell.fill.fgColor.rgb for row in sheet.iter_rows(min_row=4, min_col=8, max_col=8) for cell in row
    }
    assert len(fills) >= 4  # colour-coded statuses
    comments = [cell for row in sheet.iter_rows() for cell in row if cell.comment]
    assert comments  # comments in cells
    texts = {
        str(c.value) for row in sheet.iter_rows(min_row=4, min_col=10, max_col=10) for c in row if c.value
    }
    assert any(t.startswith("w/c") for t in texts)  # dates typed as text


def test_f05_ac05_about_eight_percent_of_typed_statuses_disagree_with_the_source(
    workbook_path: Path, plan: Plan
) -> None:
    intended = {key: stage for key, stage, _ in oracle_rows(plan, None)}
    rows, disagree = check_workbook(workbook_path, intended)
    assert rows > 400
    assert 0.06 <= disagree / rows <= 0.10, (disagree, rows)


def test_f05_fr07_the_status_vocabulary_maps_back_to_stages() -> None:
    assert normalise_status(" In QC ") == "qc_testing"
    assert normalise_status("Awaiting UD") == "qa_release"
    assert normalise_status("made up") is None
