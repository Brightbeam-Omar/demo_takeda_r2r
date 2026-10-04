"""The deliberately messy "legacy tracker" workbook (F05-FR-07): what a site tracked in spreadsheets before.

It is built from the same plan as the databases. About 8% of the Tracker tab's typed-in statuses disagree with
what the source systems say (usually a stale, earlier stage), which is what the demo's "spreadsheet versus
system" story needs. Merged headers, colour fills, comments, stale copies and a scratch tab are on purpose.
"""

import random
from collections.abc import Mapping, Sequence
from datetime import date, timedelta
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from datagen.model import BatchPlan, LotPlan, Plan, current_stage_entry
from datagen.params import Params
from datagen.rng import stream
from datagen.stats import Stats

STAGE_ORDER = [
    "pending",
    "receipt",
    "call_off",
    "sampling",
    "qc_ship",
    "qc_testing",
    "qa_release",
    "released",
]
STATUS_TEXT: dict[str, list[str]] = {
    "pending": ["Awaiting GR", "GR reversed", "Not received"],
    "receipt": ["Received", "Rec'd - check open", "Inbound check"],
    "call_off": ["Call off", "Call-off pending", "At 3PL", "Request transfer"],
    "sampling": ["To sample", "Sampling", "Awaiting sample", "Sample req"],
    "qc_ship": ["Shipping to lab", "Sample out", "At courier"],
    "qc_testing": ["In QC", "Testing", "QC testing", "With lab"],
    "qa_release": ["With QA", "QA review", "Awaiting UD", "Release pending"],
    "released": ["Released", "Rel.", "Done", "OK to use"],
}
STATUS_STAGE = {text.casefold(): stage for stage, texts in STATUS_TEXT.items() for text in texts}
FILLS = {
    "released": "C6E0B4",
    "qa_release": "BDD7EE",
    "qc_testing": "FFE699",
    "qc_ship": "FFE699",
    "sampling": "F8CBAD",
    "call_off": "F8CBAD",
    "receipt": "EDEDED",
    "pending": "D9D9D9",
}
NOTES = (
    "chased Mon",
    "waiting on supplier CoA",
    "Pat to confirm",
    "see email",
    "moved to cold store",
    "re-labelled",
    "QA aware",
    "check with lab",
    "??",
)
SHEETS = (
    "Tracker",
    "Call-Off",
    "Sampling Plan",
    "QC Queue",
    "KPI Weekly",
    "Re-Evals",
    "Deviations",
    "Campaign Notes",
    "Lookups",
    "Old",
    "Copy of Tracker",
    "Sheet3",
)
DISAGREEMENT_SHARE = 0.08
HEADER = (
    "Material",
    "Description",
    "Batch",
    "Lot",
    "Supplier",
    "Received",
    "Location",
    "Status",
    "Since",
    "Need by",
    "Campaign",
    "Notes",
)


def normalise_status(text: str) -> str | None:
    """The stage a typed-in status stands for (None if it is not in the vocabulary)."""
    return STATUS_STAGE.get(text.strip().casefold())


def _fill(colour: str) -> PatternFill:
    return PatternFill("solid", fgColor=colour)


def _title(sheet: Worksheet, text: str, width: int) -> None:
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    sheet.cell(1, 1, text).font = Font(bold=True, size=14, color="FFFFFF")
    sheet.cell(1, 1).fill = _fill("1F4E78")
    sheet.cell(1, 1).alignment = Alignment(horizontal="center")


def _header(sheet: Worksheet, row: int, labels: tuple[str, ...]) -> None:
    for column, label in enumerate(labels, start=1):
        cell = sheet.cell(row, column, label)
        cell.font = Font(bold=True)
        cell.fill = _fill("D9E1F2")
        sheet.column_dimensions[get_column_letter(column)].width = max(12, len(label) + 4)


class Source:
    """What the workbook needs to know about each lot, resolved once."""

    def __init__(self, plan: Plan, lot_numbers: Mapping[str, str] | None) -> None:
        self.plan = plan
        self.lot_numbers = lot_numbers or {}
        self.description = {m.matnr: m.description for m in plan.world.materials}
        self.campaign: dict[str, str] = {}
        for demand in sorted(plan.demands, key=lambda d: d.requirement_date):
            if demand.closed_on is None and demand.requirement_date >= plan.calendar.today:
                self.campaign.setdefault(demand.matnr, demand.campaign)

    def number(self, lot: LotPlan) -> str:
        return self.lot_numbers.get(lot.ref, lot.ref.split("|", 2)[2])

    def row_key(self, batch: BatchPlan, lot: LotPlan) -> str:
        return f"{batch.matnr}|{batch.charg}|{self.number(lot)}"


def tracker_rows(plan: Plan) -> list[tuple[BatchPlan, LotPlan]]:
    """Open lots and the lots released in the last 8 weeks: what a planner would still keep on the sheet."""
    cutoff = plan.calendar.today - timedelta(days=56)
    rows = [
        (b, lot)
        for b, lot in plan.lots()
        if lot.stage != "released" or (lot.ud_date is not None and lot.ud_date >= cutoff)
    ]
    return sorted(rows, key=lambda r: (r[0].matnr, r[0].charg, r[1].start))


def _typed_status(rng: random.Random, intended: str, disagree: bool) -> str:
    if not disagree:
        return rng.choice(STATUS_TEXT[intended])
    index = STAGE_ORDER.index(intended)
    if index > 1 and rng.random() < 0.7:
        wrong = STAGE_ORDER[index - 1]  # the sheet was not updated after the last move
    else:
        wrong = rng.choice([s for s in STAGE_ORDER if s != intended])
    return rng.choice(STATUS_TEXT[wrong])


def _tracker_sheet(
    sheet: Worksheet, source: Source, rng: random.Random, rows: list[tuple[BatchPlan, LotPlan]]
) -> None:
    _title(
        sheet,
        "INCOMING MATERIALS TRACKER - Site A - do not edit columns H-J without telling Pat",
        len(HEADER),
    )
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=5)
    sheet.merge_cells(start_row=2, start_column=6, end_row=2, end_column=9)
    sheet.cell(2, 1, "MATERIAL").font = Font(bold=True)
    sheet.cell(2, 6, "STATUS (typed by hand)").font = Font(bold=True)
    _header(sheet, 3, HEADER)
    wrong = set(rng.sample(range(len(rows)), round(DISAGREEMENT_SHARE * len(rows))))
    row_number = 4
    for index, (batch, lot) in enumerate(rows):
        if index and index % 45 == 0:
            row_number += 1  # a blank spacer row, as people leave
        status = _typed_status(rng, lot.stage, index in wrong)
        since = current_stage_entry(batch, lot)
        need = source.plan.need_by.get(batch.matnr)
        values: list[Any] = [
            batch.matnr,
            source.description[batch.matnr],
            batch.charg,
            source.number(lot),
            batch.lifnr,
            lot.start,
            batch.lgort,
            status,
            since,
            need
            if rng.random() > 0.15 or need is None
            else f"w/c {need - timedelta(days=need.weekday()):%d-%b}",
            source.campaign.get(batch.matnr, ""),
            rng.choice(NOTES) if rng.random() < 0.15 else "",
        ]
        for column, value in enumerate(values, start=1):
            sheet.cell(row_number, column, value)
        for column in (6, 9, 10):
            if isinstance(sheet.cell(row_number, column).value, date):
                sheet.cell(row_number, column).number_format = "dd-mmm-yy"
        shown = normalise_status(status) or lot.stage
        sheet.cell(row_number, 8).fill = _fill(FILLS[shown])
        if rng.random() < 0.06:
            sheet.cell(row_number, 8).comment = Comment("updated from memory, not from the ERP", "Pat")
        row_number += 1
    sheet.freeze_panes = "A4"


def _list_sheet(sheet: Worksheet, title: str, header: tuple[str, ...], rows: Sequence[Sequence[Any]]) -> None:
    _title(sheet, title, len(header))
    _header(sheet, 2, header)
    for number, row in enumerate(rows, start=3):
        for column, value in enumerate(row, start=1):
            sheet.cell(number, column, value)


def build_workbook(
    plan: Plan,
    params: Params,
    stats: Stats,
    lot_numbers: Mapping[str, str] | None = None,
) -> Workbook:
    rng = stream(plan.seed, "legacy-workbook")
    source = Source(plan, lot_numbers)
    rows = tracker_rows(plan)
    workbook = Workbook()
    first = workbook.active
    assert first is not None
    first.title = SHEETS[0]
    sheets = {name: workbook.create_sheet(name) for name in SHEETS[1:]}
    sheets[SHEETS[0]] = first
    _tracker_sheet(first, source, rng, rows)

    def lots_in(*stages: str) -> list[tuple[BatchPlan, LotPlan]]:
        return [(b, lot) for b, lot in rows if lot.stage in stages]

    _list_sheet(
        sheets["Call-Off"],
        "3PL call-off list",
        ("Material", "Batch", "3PL", "Received", "Called off?", "Chased"),
        [
            [
                b.matnr,
                b.charg,
                "3PL North" if b.lgort == "0200" else "3PL South",
                lot.start,
                "N",
                rng.choice(["", "Mon", "Wed"]),
            ]
            for b, lot in lots_in("call_off")
        ],
    )
    _list_sheet(
        sheets["Sampling Plan"],
        "Sampling plan (week by week)",
        ("Material", "Batch", "Planned date", "Sampler", "Done"),
        [
            [
                b.matnr,
                b.charg,
                plan.calendar.today + timedelta(days=rng.randint(0, 9)),
                rng.choice(["Sam", "Jo", "Lee"]),
                "",
            ]
            for b, _ in lots_in("sampling")
        ],
    )
    _list_sheet(
        sheets["QC Queue"],
        "QC queue",
        ("Material", "Batch", "Analyst", "Lab", "Priority"),
        [
            [
                b.matnr,
                b.charg,
                rng.choice(["Quinn", "Ana", "Raj"]),
                "onsite" if not (lot.latest and lot.latest.offsite) else "external",
                rng.choice(["", "HIGH", "rush"]),
            ]
            for b, lot in lots_in("qc_testing", "qc_ship")
        ],
    )
    kpi = sheets["KPI Weekly"]
    _list_sheet(
        kpi,
        "KPI weekly - typed in on Mondays",
        ("Week", "Sampling %", "Testing %", "QA release %", "Comment"),
        [],
    )
    for number, week in enumerate(plan.calendar.metric_week_starts, start=3):
        values: list[Any] = [week]
        for metric in ("M3", "M6", "M7"):
            stat = next(w for w in stats.weekly[metric] if w.week == week)
            values.append(f"{stat.pct + rng.choice([-2, 0, 0, 3]):.0f}%" if stat.pct is not None else "n/a")
        values.append(rng.choice(["", "", "holiday week", "#REF!", "lab down 2 days"]))
        for column, value in enumerate(values, start=1):
            kpi.cell(number, column, value)
    _list_sheet(
        sheets["Re-Evals"],
        "Re-evaluations due",
        ("Material", "Batch", "Lot", "Opened", "Status"),
        [
            [b.matnr, b.charg, source.number(lot), lot.start, rng.choice(STATUS_TEXT[lot.stage])]
            for b, lot in rows
            if lot.lot_type == "09"
        ],
    )
    _list_sheet(
        sheets["Deviations"],
        "Deviations (QA sends this on Fridays)",
        ("No.", "Title", "Severity", "Opened", "Open?"),
        [
            [f"D-{n:03d}", d.title, d.severity, d.opened_on, "Y" if d.closed_on is None else "N"]
            for n, d in enumerate(plan.deviations, start=1)
        ],
    )
    _list_sheet(
        sheets["Campaign Notes"],
        "Campaign notes",
        ("Campaign", "Note"),
        [
            [
                c,
                rng.choice(
                    [
                        "watch the buffer salts",
                        "needs API intermediates by month end",
                        "dates keep moving",
                        "ask planning",
                    ]
                ),
            ]
            for c in plan.world.campaigns
        ],
    )
    _list_sheet(
        sheets["Lookups"],
        "Lookups (do not delete)",
        ("Status", "Stage it means"),
        [[text, stage] for stage, texts in STATUS_TEXT.items() for text in texts],
    )
    _list_sheet(
        sheets["Old"],
        "OLD - do not use",
        ("Material", "Batch", "Status"),
        [
            [b.matnr, b.charg, "Released"]
            for b, lot in plan.lots()
            if lot.stage == "released" and lot not in [r[1] for r in rows]
        ][:80],
    )
    stale = sheets["Copy of Tracker"]
    _list_sheet(
        stale,
        "Copy of Tracker (v2) - last month",
        ("Material", "Batch", "Status", "Notes"),
        [[b.matnr, b.charg, rng.choice(STATUS_TEXT[rng.choice(STAGE_ORDER)]), "old"] for b, _ in rows[:60]],
    )
    scratch = sheets["Sheet3"]
    for number in range(1, 15):
        scratch.cell(number, 1, rng.randint(1, 400))
        scratch.cell(number, 2, rng.choice(["???", "tmp", "check", ""]))
    scratch.cell(16, 1, "=SUM(A1:A14)")
    for sheet in workbook.worksheets:
        sheet.sheet_properties.tabColor = rng.choice(["FF0000", "00B050", "FFC000", "5B9BD5"])
    return workbook


def check_workbook(path: Path, intended: Mapping[str, str]) -> tuple[int, int]:
    """``(tracker rows, rows whose typed status disagrees with the intended stage)`` of a saved workbook."""
    sheet = load_workbook(path)["Tracker"]
    rows = disagree = 0
    for values in sheet.iter_rows(min_row=4, values_only=True):
        material, batch, lot, status = values[0], values[2], values[3], values[7]
        if not material or status is None:
            continue
        key = f"{material}|{batch}|{lot}"
        if key not in intended:
            continue
        rows += 1
        disagree += int(normalise_status(str(status)) != intended[key])
    return rows, disagree
