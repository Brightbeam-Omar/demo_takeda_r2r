"""T3: xlsx / office / csv / json / parquet / Delta readers [F02-FR-04, F02-AC-02]."""

import zipfile
from pathlib import Path

import openpyxl
import pyarrow as pa
import pyarrow.parquet as pq
import pytest
from deltalake import write_deltalake
from leakscan.load import parse_denylist
from leakscan.scan import SkippedFile, scan_file, scan_path

RULES = parse_denylist("foobarco")


def test_f02_ac02_xlsx_cell_is_found_with_sheet_and_cell_reference(tmp_path: Path) -> None:
    """F02-AC-02: a term inside an .xlsx cell fails with a sheet!cell reference."""
    book = openpyxl.Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Plan"
    sheet["A1"] = "header"
    sheet["B3"] = "Visit to FooBarCo site"
    sheet["C3"] = 42
    book.save(tmp_path / "plan.xlsx")
    [finding] = scan_file(tmp_path / "plan.xlsx", "plan.xlsx", RULES)
    assert finding.render() == "plan.xlsx:Plan!B3: F*** (rule #1)"


def test_f02_fr04_xlsx_document_properties_are_scanned(tmp_path: Path) -> None:
    book = openpyxl.Workbook()
    book.properties.creator = "Someone at FooBarCo"
    book.save(tmp_path / "p.xlsx")
    assert [f.location for f in scan_file(tmp_path / "p.xlsx", "p.xlsx", RULES)] == ["properties"]


def _office(path: Path, parts: dict[str, str]) -> None:
    with zipfile.ZipFile(path, "w") as archive:
        for name, body in parts.items():
            archive.writestr(name, body)


def test_f02_oq014_docx_text_split_across_runs_is_found(tmp_path: Path) -> None:
    body = (
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>'
        "<w:p><w:r><w:t>Foo</w:t></w:r><w:r><w:t>BarCo</w:t></w:r></w:p><w:p><w:r><w:t>clean</w:t></w:r></w:p>"
        "</w:body></w:document>"
    )
    _office(tmp_path / "memo.docx", {"word/document.xml": body})
    [finding] = scan_file(tmp_path / "memo.docx", "memo.docx", RULES)
    assert finding.location == "word/document.xml"


def test_f02_oq014_pptx_slide_text_is_found(tmp_path: Path) -> None:
    body = (
        '<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><p:cSld><p:spTree><p:sp><p:txBody>'
        "<a:p><a:r><a:t>FooBarCo roadmap</a:t></a:r></a:p></p:txBody></p:sp></p:spTree></p:cSld></p:sld>"
    )
    _office(tmp_path / "deck.pptx", {"ppt/slides/slide1.xml": body})
    assert len(scan_file(tmp_path / "deck.pptx", "deck.pptx", RULES)) == 1


def test_f02_oq014_office_core_properties_are_scanned(tmp_path: Path) -> None:
    core = (
        "<cp:coreProperties xmlns:cp='x' xmlns:dc='y'><dc:creator>FooBarCo</dc:creator></cp:coreProperties>"
    )
    _office(tmp_path / "memo.docx", {"docProps/core.xml": core})
    assert len(scan_file(tmp_path / "memo.docx", "memo.docx", RULES)) == 1


@pytest.mark.parametrize(
    ("name", "content"),
    [("data.csv", "id,site\n1,FooBarCo\n"), ("data.json", '{"site": "FooBarCo"}\n')],
)
def test_f02_fr04_csv_and_json_are_scanned_as_text(tmp_path: Path, name: str, content: str) -> None:
    (tmp_path / name).write_text(content)
    assert len(scan_file(tmp_path / name, name, RULES)) == 1


def test_f02_fr04_parquet_column_names_and_string_rows_are_scanned(tmp_path: Path) -> None:
    table = pa.table({"foobarco_id": [1, 2], "site": ["ok", "FooBarCo plant"], "qty": [7, 8]})
    pq.write_table(table, tmp_path / "t.parquet")
    findings = scan_file(tmp_path / "t.parquet", "t.parquet", RULES)
    assert sorted(f.location for f in findings) == ["column#1", "r2c2"]


def test_f02_oq013_every_row_is_scanned_with_no_cap(tmp_path: Path) -> None:
    values = ["clean"] * 5000 + ["FooBarCo"]
    pq.write_table(pa.table({"site": values}), tmp_path / "big.parquet")
    [finding] = scan_file(tmp_path / "big.parquet", "big.parquet", RULES)
    assert finding.location == "r5001c1"


def test_f02_fr04_delta_table_data_and_log_are_covered(tmp_path: Path) -> None:
    """Delta tables are parquet data files plus a JSON log; both reach a reader."""
    table_dir = tmp_path / "fixtures" / "tbl"
    write_deltalake(str(table_dir), pa.table({"site": ["FooBarCo plant", "other"]}))
    hits = [
        f
        for p in sorted(table_dir.rglob("*"))
        if p.is_file() and not p.name.endswith(".crc")
        for f in scan_file(p, p.relative_to(tmp_path).as_posix(), RULES)
    ]
    assert any(f.path.endswith(".parquet") and f.location == "r1c1" for f in hits)


def test_f02_oq014_other_binaries_are_skipped_not_scanned(tmp_path: Path) -> None:
    (tmp_path / "scan.pdf").write_bytes(b"%PDF-1.4\x00\x01FooBarCo")
    with pytest.raises(SkippedFile, match="binary"):
        scan_file(tmp_path / "scan.pdf", "scan.pdf", RULES)


def test_f02_oq014_unreadable_office_file_is_skipped_with_reason(tmp_path: Path) -> None:
    (tmp_path / "broken.xlsx").write_bytes(b"not a zip")
    with pytest.raises(SkippedFile, match="unreadable"):
        scan_file(tmp_path / "broken.xlsx", "broken.xlsx", RULES)


def test_f02_oq014_file_paths_are_scanned(tmp_path: Path) -> None:
    [finding] = scan_path("clients/FooBarCo/readme.txt", RULES)
    assert finding.render() == "clients/FooBarCo/readme.txt:(path): F*** (rule #1)"


def test_f02_fr05_locations_that_match_a_rule_are_masked(tmp_path: Path) -> None:
    """A sheet called after the term must not leak through the location."""
    book = openpyxl.Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "FooBarCo"
    sheet["A1"] = "FooBarCo"
    book.save(tmp_path / "s.xlsx")
    findings = scan_file(tmp_path / "s.xlsx", "s.xlsx", RULES)
    assert len(findings) == 2  # the sheet name and the cell
    assert all("foobarco" not in f.render().lower() for f in findings)
